"""Create the final formal Word report from reviewed Markdown; renderer QA is a separate required step."""
import argparse,re
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
p=argparse.ArgumentParser();p.add_argument('--input',default='reports/FINAL_RESULTS.md');p.add_argument('--output',default='reports/trusted_agent_reactions_report.docx');a=p.parse_args()
source=Path(a.input)
if not source.exists():raise SystemExit('Reviewed final Markdown report is required')
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.left_margin=sec.right_margin=Inches(.75);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65)
for name,size in [('Normal',11),('Title',28),('Heading 1',18),('Heading 2',12.5),('Heading 3',11)]:
 s=doc.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(7 if name=='Normal'else 9);s.paragraph_format.line_spacing=1.1
 if name.startswith('Heading'):s.paragraph_format.space_before=Pt(15);s.paragraph_format.keep_with_next=True
for name in ['Title','Subtitle']:
 s=doc.styles[name]
 if s.element.pPr is not None:
  for child in list(s.element.pPr):
   if child.tag==qn('w:pBdr'):s.element.pPr.remove(child)
 s.font.color.rgb=RGBColor(0,0,0)

def clean(t):return t.replace('−','-').replace('–','-').replace('—','-').replace('→','to').replace('×','x')
def inline(p,text):
 for piece in re.split(r'(\*\*.*?\*\*)',clean(text)):
  bold=piece.startswith('**')and piece.endswith('**');piece=piece[2:-2]if bold else piece
  piece=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',piece).replace('`','')
  r=p.add_run(piece);r.bold=bold

def table(lines):
 rows=[[x.strip()for x in line.strip().strip('|').split('|')]for line in lines]
 rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',v.replace(' ',''))for v in r)]
 if not rows:return
 n=len(rows[0]);widths={2:[2.0,5.0],3:[1.6,2.7,2.7],4:[1.75,1.9,1.9,1.45],5:[1.7,1.5,1.5,.9,1.4]}.get(n,[7/n]*n)
 t=doc.add_table(rows=1,cols=n);t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 props=t._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 props.append(borders)
 for i,values in enumerate(rows):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  tr=t.rows[i]._tr;trPr=tr.get_or_add_trPr();no_split=OxmlElement('w:cantSplit');trPr.append(no_split)
  if i==0:
   repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for j,cell in enumerate(cells):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for side,value in [('top','95'),('bottom','95'),('left','110'),('right','110')]:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');m.append(e)
   cp.append(m);fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'203B50'if i==0 else'EFF4F7'if i%2==0 else'FFFFFF');cp.append(fill)
   para=cell.paragraphs[0];para.paragraph_format.space_after=Pt(0);para.paragraph_format.space_before=Pt(0);para.paragraph_format.line_spacing=1.05
   inline(para,values[j]if j<len(values)else'')
   for run in para.runs:run.font.size=Pt(9.5);run.font.color.rgb=RGBColor(255,255,255)if i==0 else RGBColor(0,0,0);run.bold=i==0 or run.bold
   if i==0:para.paragraph_format.keep_with_next=True
 doc.add_paragraph().paragraph_format.space_after=Pt(2)

lines=source.read_text().splitlines();i=0;first=True
while i<len(lines):
 line=lines[i].strip()
 if not line:i+=1;continue
 if line.startswith('|'):
  group=[]
  while i<len(lines)and lines[i].lstrip().startswith('|'):group.append(lines[i]);i+=1
  table(group);continue
 if line.startswith('#'):
  level=len(line)-len(line.lstrip('#'));text=line[level:].strip()
  # Keep headings natural and unornamented, following the reference document style.
  if first:
   inline(doc.add_paragraph(style='Title'),text);first=False
   sub=doc.add_paragraph('Expanded local market experiment');sub.paragraph_format.space_after=Pt(16)
  else:
   heading=doc.add_paragraph(style='Heading '+str(min(level,3)));inline(heading,text)
 elif line.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),line[2:])
 else:inline(doc.add_paragraph(),line)
 i+=1
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Page ');r.font.size=Pt(9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.title='Colluding firms reactions to a trusted peer';doc.core_properties.subject='Expanded market experiment design and observed results';doc.core_properties.author='';doc.core_properties.last_modified_by=''
output=Path(a.output);output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
print('Saved',output)
